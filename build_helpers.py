"""
Complete Academic Project Report Generator for:
AUTHENTICATED FILE ENCRYPTION SYSTEM (AFES)
Generates:
1. Authenticated_File_Encryption_System_Project_Report.docx
2. Authenticated_File_Encryption_System_Project_Report.pdf (via MS Word COM)
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_NAME = "Authenticated_File_Encryption_System_Project_Report.docx"
PDF_NAME = "Authenticated_File_Encryption_System_Project_Report.pdf"

# Color Palette matching reference template
C_NAVY = RGBColor(30, 58, 138)       # #1e3a8a
C_DARK_SLATE = RGBColor(15, 23, 42)  # #0f172a
C_BODY = RGBColor(31, 41, 55)        # #1f2937
C_MUTED = RGBColor(148, 163, 184)    # #94a3b8
C_EMERALD = RGBColor(5, 150, 105)    # #059669
C_AMBER = RGBColor(217, 119, 6)      # #d97706
C_BLUE = RGBColor(37, 99, 235)       # #2563eb

HEX_HEADER_BG = "0F172A"             # Dark navy table header
HEX_ZEBRA_BG = "F8FAFC"              # Very light grey zebra row
HEX_BORDER = "CBD5E1"                # Slate-300 table borders
HEX_CALLOUT_BG = "F0FDF4"            # Emerald tint
HEX_CALLOUT_BORDER = "059669"        # Emerald left border
HEX_CODE_BG = "F8FAFC"               # Code block background

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding in twips (1/20 of a pt)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set cell background color."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color=HEX_BORDER, sz="4", val="single"):
    """Set subtle grey borders on a table."""
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="none"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_callout_box(doc, title, lines, border_hex="2563eb", bg_hex="eff6ff"):
    """Adds a stylish callout box matching the reference report."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=120, bottom=120, left=200, right=200)
    set_cell_shading(cell, bg_hex)
    
    # Left border only
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    '''
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title.upper()}]\n")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = C_NAVY
    
    for l_idx, line in enumerate(lines):
        p_line = cell.add_paragraph()
        p_line.paragraph_format.space_before = Pt(1)
        p_line.paragraph_format.space_after = Pt(1)
        p_line.paragraph_format.line_spacing = 1.15
        r = p_line.add_run(line)
        r.font.name = "Calibri"
        r.font.size = Pt(9.0)
        r.font.color.rgb = C_BODY
        if line.startswith("•") or line.startswith("Where:"):
            p_line.paragraph_format.left_indent = Inches(0.15)
    
    # Spacing after callout box
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_code_listing(doc, caption, code_lines):
    """Add a shaded code block with Consolas font."""
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(6)
    p_cap.paragraph_format.space_after = Pt(2)
    r_cap = p_cap.add_run(caption)
    r_cap.bold = True
    r_cap.font.name = "Calibri"
    r_cap.font.size = Pt(9.5)
    r_cap.font.color.rgb = C_NAVY

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_margins(cell, top=100, bottom=100, left=180, right=180)
    set_cell_shading(cell, HEX_CODE_BG)
    
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>
        <w:left w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>
        <w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>
    </w:tcBorders>
    '''
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    for idx, c_line in enumerate(code_lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
        r = p.add_run(c_line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = C_DARK_SLATE

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

print("Helper definitions loaded successfully.")
